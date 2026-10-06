"""Extrae los precios al público de las planillas mensuales de Posventa ANZER.

Uso:
    python extraer_precios.py Detalle_Servicios.xlsx LISTA_DE_PRECIOS_POSVENTA_ANZER.xlsx salida.json

La primera planilla aporta los services oficiales Ford. La segunda aporta el
Service Promo (vehículos fuera de garantía) y las distribuciones, que se
calculan con los precios con IVA (columna D) de la pestaña "Lista de Precios FORD".

Solo se leen precios sugeridos al público (IVA incluido). Los costos internos
(dealer net, márgenes) de la planilla nunca se exportan.
"""
import json
import re
import sys

import openpyxl

HOJAS_SERVICE = ["New Portfolio", "Old Portfolio", "Precios Vigentes Electricos"]
CATEGORIAS = {"AUTOS", "SUV´S", "PICKS-UPS", "PERFORMANCE", "COMERCIALES", "UTILITARIOS"}


def num(v):
    return round(v) if isinstance(v, (int, float)) else None


def limpiar(s):
    return re.sub(r"\s+", " ", str(s)).strip(" -")


def services_programados(wb):
    modelos = []
    for nombre in HOJAS_SERVICE:
        ws = wb[nombre]
        categoria, kms = None, []
        for row in ws.iter_rows(values_only=True):
            celdas = list(row[1:])
            if celdas and isinstance(celdas[0], str) and celdas[0].strip() in CATEGORIAS:
                categoria = celdas[0].strip()
                continue
            numeros = [c for c in celdas[2:] if isinstance(c, (int, float))]
            if celdas[0] in (None, "") and len(numeros) >= 5 and all(n < 1000 for n in numeros):
                kms = [int(n * 1000) for n in numeros]  # fila "Kilómetros (miles)"
                continue
            if len(celdas) > 1 and celdas[1] == "Fidelidad Ford" and celdas[0]:
                precios = [num(c) for c in celdas[2 : 2 + len(kms)]]
                modelos.append({
                    "modelo": limpiar(celdas[0]),
                    "categoria": categoria,
                    "hoja": nombre,
                    "services": [
                        {"km": km, "precio": p} for km, p in zip(kms, precios) if p
                    ],
                })
    return modelos


def mano_de_obra_y_lavado(wb):
    ws = wb["Mano de Obra"]
    hora, lavado, seccion = {}, {}, None
    for row in ws.iter_rows(values_only=True):
        c = [x for x in row if x not in (None, "")]
        if not c:
            continue
        if isinstance(c[0], str) and "PAUTA LAVADO" in str(c):
            seccion = "lavado"
            continue
        if isinstance(c[0], str) and isinstance(c[-1], (int, float)) and c[-1] > 1000:
            destino = lavado if seccion == "lavado" else hora
            destino[limpiar(c[0])] = num(c[-1])
    return {"valor_hora_con_iva": hora, "lavado_con_iva": lavado}


def combos(wb):
    ws = wb["Oport Mecanica Lig"]
    resultado, accion, columnas = {}, None, []
    for row in ws.iter_rows(values_only=True):
        c = list(row[1:])
        if isinstance(c[0], str) and c[0].startswith("Accion"):
            accion = limpiar(c[0])
            columnas = [limpiar(x) for x in c[1:] if x]
            resultado[accion] = []
        elif accion and isinstance(c[0], str):
            valores = [num(x) for x in c[1:] if x not in (None, "")]
            if any(valores) and not str(c[0]).startswith("Distir"):
                resultado[accion].append({"modelo": limpiar(c[0]), **dict(zip(columnas, valores))})
            elif str(c[0]).startswith("Distir"):
                accion = None
    return resultado


def redondear_100(x):
    """Igual que MROUND(x, 100) de Excel."""
    return int(x / 100 + 0.5) * 100


def lista_ford(wb):
    """Código de pieza -> precio público con IVA (columna D)."""
    precios = {}
    for row in wb["Lista de Precios FORD"].iter_rows(values_only=True):
        if row[0] and isinstance(row[3], (int, float)):
            precios[str(row[0]).strip().upper()] = row[3]
    return precios


def precio_pieza(precios, codigo, obligatoria=True):
    """Precio con IVA de una pieza. Una pieza opcional que falta en la lista vale None."""
    codigo = str(codigo or "").strip().upper()
    if codigo in ("", "-"):
        return 0
    if codigo not in precios:
        if obligatoria:
            raise KeyError(f"La pieza {codigo} no está en la Lista de Precios FORD")
        return None
    return precios[codigo]


def service_promo(wb, precios):
    ws = wb["SERVIS PROMO "]
    aceite = {str(ws[f"B{f}"].value).strip().upper(): ws[f"C{f}"].value for f in (25, 26)}
    mano_de_obra = ws["G34"].value
    opcionales = {"filtro_aire_motor": "G", "filtro_combustible": "H", "filtro_habitaculo": "I"}
    modelos = []
    for f in range(2, 24):
        nombre = ws[f"B{f}"].value
        if not nombre:
            continue
        litros = ws[f"C{f}"].value
        repuestos = litros * aceite[str(ws[f"D{f}"].value).strip().upper()] + precio_pieza(precios, ws[f"F{f}"].value)
        precio = redondear_100(repuestos) + mano_de_obra
        modelos.append({
            "modelo": limpiar(nombre),
            "litros_aceite": litros,
            "viscosidad": str(ws[f"D{f}"].value).strip().upper(),
            "precio_servicio_minimo": precio,
            "precio_con_descuento_contado": redondear_100(precio * 0.9),
            "opcionales": {  # None = no figura en la lista, se informa "a consultar"
                k: (lambda p: round(p) if p is not None else None)(precio_pieza(precios, ws[f"{col}{f}"].value, obligatoria=False))
                for k, col in opcionales.items()
            },
        })
    return {
        "solo_fuera_de_garantia": True,
        "incluye": ["Aceite", "Filtro de aceite", "Lavado de cortesía", "Revisión general del vehículo"],
        "mano_de_obra": mano_de_obra,
        "modelos": modelos,
    }


def distribuciones(wb, precios):
    ws = wb["DISTRIBUCIONES"]
    hora = ws["G19"].value
    litro_aceite = wb["SERVIS PROMO "]["C25"].value  # 5W30
    reten_ciguenal = ws["D25"].value  # lo llevan todos los motores
    motores = []
    for f in range(2, 15):
        nombre = ws[f"B{f}"].value
        if not nombre:
            continue
        litros = ws[f"C{f}"].value
        repuestos = (
            (litros * litro_aceite if isinstance(litros, (int, float)) else 0)
            # correa, tensor, bomba de agua, junta de bomba, correa poli V, correa bomba de aceite, filtro
            + sum(precio_pieza(precios, ws[f"{col}{f}"].value) for col in "EFGHIJK")
            + reten_ciguenal
            + precio_pieza(precios, ws[f"M{f}"].value) * 2  # 2 bidones de refrigerante
        )
        mano_de_obra = ws[f"N{f}"].value * hora
        precio = redondear_100(repuestos + 20) + mano_de_obra
        motores.append({
            "motor": limpiar(nombre),
            "horas_mano_de_obra": ws[f"N{f}"].value,
            "precio": precio,
            "precio_con_descuento_contado": redondear_100(precio * 0.9),
        })
    return {"valor_hora_mano_de_obra": hora, "motores": motores}


def main():
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    lista = openpyxl.load_workbook(sys.argv[2], data_only=True)
    precios = lista_ford(lista)
    datos = {
        "fuente": "LISTA DE PRECIOS POSVENTA ANZER",
        "notas": [
            "Precios sugeridos al público, IVA incluido.",
            "Precio de contado: efectivo, débito, crédito en 1 pago u otros medios electrónicos.",
            "Services: incluyen aceite sintético Elaion, filtros y operaciones de control según el manual.",
            "No incluyen líquido de frenos/embrague ni refrigerante salvo que se indique.",
        ],
        "services_programados": services_programados(wb),
        "mano_de_obra": mano_de_obra_y_lavado(wb),
        "combos_del_mes": combos(wb),
        "service_promo": service_promo(lista, precios),
        "distribuciones": distribuciones(lista, precios),
        "alineacion_y_balanceo": 55000,
    }
    with open(sys.argv[3], "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=2)
    print(f"{len(datos['services_programados'])} modelos exportados")


if __name__ == "__main__":
    main()
